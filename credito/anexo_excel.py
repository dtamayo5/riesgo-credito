"""
Anexo auditable en Excel.

Los resultados de modelos (PD, LGD, ECL por crédito) se escriben como valores;
la aritmética que el auditor quiere rehacer (totales por etapa, provisión de
referencia, base de préstamo, excedente financiero) va como fórmulas vivas que
apuntan a la hoja de detalle por crédito.
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import niif9
from .graficas import etiqueta_cosecha

ENC_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=9)
ENC_FILL = PatternFill("solid", fgColor="1F3B5C")
FUENTE = Font(name="Arial", size=9)
NEGRITA = Font(name="Arial", size=9, bold=True)
TITULO = Font(name="Arial", size=12, bold=True, color="1F3B5C")
COLOR_ESTADO = {"ROJO": "D03B3B", "ÁMBAR": "FAB219", "VERDE": "0CA30C"}
F_MILES, F_PCT, F_DEC = "#,##0", "0.0%", "0.00"


def _titulo(ws, texto, nota=None):
    ws["A1"] = texto
    ws["A1"].font = TITULO
    if nota:
        ws["A2"] = nota
        ws["A2"].font = Font(name="Arial", size=8, italic=True, color="52514E")


def _tabla(ws, df: pd.DataFrame, fila: int, col: int = 1, formatos: dict | None = None, anchos: dict | None = None):
    formatos = formatos or {}
    for j, c in enumerate(df.columns, start=col):
        cel = ws.cell(fila, j, str(c))
        cel.font, cel.fill = ENC_FONT, ENC_FILL
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, (_, r) in enumerate(df.iterrows(), start=fila + 1):
        for j, c in enumerate(df.columns, start=col):
            v = r[c]
            if isinstance(v, (np.floating, float)) and (np.isnan(v) or np.isinf(v)):
                v = None
            elif isinstance(v, np.integer):
                v = int(v)
            elif isinstance(v, np.floating):
                v = float(v)
            elif isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            cel = ws.cell(i, j, v)
            cel.font = FUENTE
            if c in formatos:
                cel.number_format = formatos[c]
    for j, c in enumerate(df.columns, start=col):
        ancho = (anchos or {}).get(c, max(10, min(40, len(str(c)) + 2)))
        ws.column_dimensions[get_column_letter(j)].width = ancho
    ws.row_dimensions[fila].height = 30
    return fila + len(df) + 1


def generar(r: dict, ruta: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    e, foto, r9, fon = r["entidad"], r["foto"], r["niif9"], r["fondeo"]
    _titulo(ws, f"{e['nombre']} - Riesgo de crédito al {r['corte']:%Y-%m-%d}",
            "Cifras en pesos salvo indicación. Demostración: datos sintéticos." if e.get("demo") else None)
    lim = r["limites"][["estado", "codigo", "indicador", "categoria", "tipo", "valor", "limite", "alerta", "unidad", "detalle"]]
    fin = _tabla(ws, lim.rename(columns=str.capitalize), 4, formatos={"Valor": "0.000", "Limite": "0.000", "Alerta": "0.000"},
                 anchos={"Indicador": 60, "Detalle": 18, "Estado": 10, "Codigo": 20})
    for i, est in enumerate(lim["estado"], start=5):
        if est in COLOR_ESTADO:
            ws.cell(i, 1).fill = PatternFill("solid", fgColor=COLOR_ESTADO[est])
    acc = pd.DataFrame(r["acciones"]).rename(columns={"prioridad": "Prioridad", "accion": "Acción", "monto": "Monto (COP)",
                                                       "sustento": "Sustento"})
    ws.cell(fin + 1, 1, "Acciones sugeridas").font = NEGRITA
    if len(acc):
        _tabla(ws, acc, fin + 2, formatos={"Monto (COP)": F_MILES}, anchos={"Acción": 60, "Sustento": 60})

    # ---- detalle por crédito (base de las fórmulas)
    det = r9.detalle
    cols = {"id": "Crédito", "producto": "Producto", "saldo_capital": "Saldo", "dias_mora": "Días de mora",
            "tramo": "Tramo", "mob": "MOB", "plazo_restante": "Plazo restante", "tasa_ea": "Tasa EA",
            "score": "Score", "calificacion": "Calificación", "pd_orig": "PD originación", "pd12": "PD 12m",
            "pd_vida": "PD vida", "lgd": "LGD", "curado": "Curado 6m", "etapa": "Etapa", "ecl": "ECL ponderada",
            "categoria": "Categoría ref.", "porc_ref": "% ref."}
    d = det[list(cols)].rename(columns=cols).copy()
    d["Curado 6m"] = d["Curado 6m"].map({True: "Sí", False: "No"})
    wd = wb.create_sheet("Creditos_NIIF9")
    _titulo(wd, "Detalle por crédito vigente", "Provisión ref. = Saldo x % ref. (fórmula). ECL: resultado del modelo.")
    _tabla(wd, d, 4, formatos={"Saldo": F_MILES, "Tasa EA": F_PCT, "PD originación": "0.00%", "PD 12m": "0.00%",
                               "PD vida": "0.00%", "LGD": F_PCT, "ECL ponderada": F_MILES, "% ref.": F_PCT})
    n = len(d)
    col_ref = len(cols) + 1
    wd.cell(4, col_ref, "Provisión ref.").font = ENC_FONT
    wd.cell(4, col_ref).fill = ENC_FILL
    for i in range(5, 5 + n):
        c = wd.cell(i, col_ref, f"=C{i}*S{i}")
        c.number_format, c.font = F_MILES, FUENTE
    wd.column_dimensions[get_column_letter(col_ref)].width = 14
    wd.freeze_panes = "B5"
    rng = lambda letra: f"Creditos_NIIF9!${letra}$5:${letra}${4 + n}"  # noqa: E731

    # ---- provisiones con fórmulas
    wp = wb.create_sheet("Provisiones")
    _titulo(wp, "Provisiones NIIF 9 y referencia regulatoria", "Fórmulas sobre la hoja Creditos_NIIF9.")
    enc = ["Etapa", "Créditos", "Saldo", "ECL", "Cobertura"]
    for j, c in enumerate(enc, start=1):
        x = wp.cell(4, j, c)
        x.font, x.fill = ENC_FONT, ENC_FILL
    for k, et in enumerate([1, 2, 3], start=5):
        wp.cell(k, 1, et).font = FUENTE
        wp.cell(k, 2, f"=COUNTIF({rng('P')},A{k})")
        wp.cell(k, 3, f"=SUMIF({rng('P')},A{k},{rng('C')})").number_format = F_MILES
        wp.cell(k, 4, f"=SUMIF({rng('P')},A{k},{rng('Q')})").number_format = F_MILES
        wp.cell(k, 5, f"=IFERROR(D{k}/C{k},0)").number_format = F_PCT
    wp.cell(8, 1, "Total").font = NEGRITA
    for j, L in [(2, "B"), (3, "C"), (4, "D")]:
        wp.cell(8, j, f"=SUM({L}5:{L}7)").number_format = F_MILES
    wp.cell(8, 5, "=D8/C8").number_format = F_PCT
    filas = [("Provisión de referencia individual", f"=SUM({rng(get_column_letter(col_ref))})"),
             ("Provisión general (% de la cartera)", e["prov_general"]),
             ("Provisión general", "=C8*B11"), ("Referencia regulatoria total", "=B10+B12"),
             ("Provisión registrada", None if np.isnan(e["prov_registrada"]) else e["prov_registrada"]),
             ("Diferencia NIIF 9 - registrada", "=D8-B14"), ("Diferencia referencia - registrada", "=B13-B14")]
    for k, (t, v) in enumerate(filas, start=10):
        wp.cell(k, 1, t).font = FUENTE
        c = wp.cell(k, 2, v)
        c.number_format = F_PCT if "(%" in t else F_MILES
    wp.column_dimensions["A"].width = 38
    for L in "BCDE":
        wp.column_dimensions[L].width = 16
    fila = 19
    wp.cell(fila, 1, "Por producto").font = NEGRITA
    pp = niif9.por_producto(r9).reset_index().rename(columns={"producto": "Producto", "saldo": "Saldo", "ecl": "ECL",
                                                              "ref": "Ref. individual", "ecl_e1": "ECL etapa 1",
                                                              "ecl_e2": "ECL etapa 2", "ecl_e3": "ECL etapa 3",
                                                              "cobertura": "Cobertura", "cobertura_ref": "Cobertura ref."})
    fila = _tabla(wp, pp, fila + 1, formatos={c: F_MILES for c in pp.columns[1:7]} | {"Cobertura": F_PCT,
                                                                                    "Cobertura ref.": F_PCT})
    wp.cell(fila + 1, 1, "Por categoría").font = NEGRITA
    pc = niif9.por_categoria(r9).reset_index()
    fila = _tabla(wp, pc, fila + 2, formatos={"saldo": F_MILES, "prov_ref": F_MILES, "ecl": F_MILES, "participacion": F_PCT})
    wp.cell(fila + 1, 1, "Escenarios").font = NEGRITA
    fila = _tabla(wp, r9.escenarios, fila + 2,
                  formatos={c: F_MILES for c in r9.escenarios.columns if c.startswith(("ecl", "saldo"))})
    wp.cell(fila + 1, 1, "Matriz de referencia aplicada (hoja Matriz_provision de la política)").font = NEGRITA
    _tabla(wp, r["datos"].matriz, fila + 2, formatos={"provision": F_PCT})

    # ---- parámetros de pérdida
    wl = wb.create_sheet("PD_LGD")
    _titulo(wl, "Parámetros de PD y LGD")
    pdc = r9.pd_cadena.reset_index().rename(columns={"index": "Producto", "0": "PD al día", "1": "PD 1-30",
                                                     "2": "PD 31-60", "3": "PD 61-90"})
    fila = _tabla(wl, pdc, 3, formatos={c: "0.00%" for c in pdc.columns[1:]})
    fila = _tabla(wl, r9.factor_pit.rename("Factor PIT").reset_index().rename(columns={"index": "Producto"}), fila + 2,
                  formatos={"Factor PIT": "0.000"})
    lg = r9.lgd.reset_index()
    fila = _tabla(wl, lg, fila + 2, formatos={"ead_incumplida": F_MILES, "recuperacion": F_PCT,
                                              "recuperacion_descontada": F_PCT, "lgd": F_PCT, "tasa_descuento": F_PCT,
                                              "lgd_usada": F_PCT})
    cr = r9.curva_recuperacion.reset_index().rename(columns={"index": "Meses desde el incumplimiento"})
    fila = _tabla(wl, cr, fila + 2, formatos={c: F_PCT for c in cr.columns[1:]})
    for prod, (mat, nobs) in r9.matrices.items():
        wl.cell(fila + 1, 1, f"Matriz de transición mensual - {'Toda la cartera' if prod == '__total__' else prod} "
                             f"({nobs} transiciones)").font = NEGRITA
        fila = _tabla(wl, mat.reset_index().rename(columns={"index": "Desde"}), fila + 2,
                      formatos={c: F_PCT for c in mat.columns})

    # ---- cartera, cosechas, rodamientos
    wc = wb.create_sheet("Cartera")
    _titulo(wc, "Cartera")
    fila = _tabla(wc, r["productos"].reset_index(), 3, formatos={"saldo": F_MILES, "participacion": F_PCT, "tasa": F_PCT,
                                                                 "icv30": F_PCT, "icv90": F_PCT})
    fila = _tabla(wc, r["tramos"].reset_index(), fila + 2, formatos={"saldo": F_MILES, "participacion": F_PCT})
    ev = r["evolucion"].reset_index()
    fila = _tabla(wc, ev, fila + 2, formatos={c: F_MILES for c in ["cartera", "saldo_30", "saldo_90", "elegible",
                                                                   "desembolsos", "castigos", "recuperaciones",
                                                                   "ticket_promedio"]} |
                  {c: F_PCT for c in ["icv30", "icv90", "tasa_promedio", "perdida_neta_12m"]} | {"mes": "yyyy-mm"})
    t20 = r["concentracion"]["top20_tabla"].reset_index().rename(columns={"id_deudor": "Deudor", "saldo_capital": "Saldo"})
    _tabla(wc, t20, fila + 2, formatos={"Saldo": F_MILES})

    wk = wb.create_sheet("Cosechas")
    _titulo(wk, "Cosechas: % acumulado de créditos que superaron 30 y 90 días, por MOB")
    c30 = r["cosechas30"].copy()
    c30.index = [etiqueta_cosecha(q) for q in c30.index]
    fila = _tabla(wk, c30.reset_index().rename(columns={"index": "Cosecha (30+)"}), 3,
                  formatos={c: F_PCT for c in c30.columns})
    c90 = r["cosechas90"].copy()
    c90.index = [etiqueta_cosecha(q) for q in c90.index]
    _tabla(wk, c90.reset_index().rename(columns={"index": "Cosecha (90+)"}), fila + 2, formatos={c: F_PCT for c in c90.columns})

    wr = wb.create_sheet("Rodamientos")
    _titulo(wr, "Matrices de rodamiento (promedio últimos meses)")
    fila = _tabla(wr, r["rodamientos"].reset_index().rename(columns={"index": "Por saldo"}), 3,
                  formatos={c: F_PCT for c in r["rodamientos"].columns})
    _tabla(wr, r["rodamientos_n"].reset_index().rename(columns={"index": "Por número"}), fila + 2,
           formatos={c: F_PCT for c in r["rodamientos_n"].columns})

    # ---- score
    sc = r["score"]
    wsx = wb.create_sheet("Score")
    _titulo(wsx, "Score de originación")
    met = pd.DataFrame({k: v for k, v in sc.metricas.items() if isinstance(v, dict)}).T.reset_index()
    fila = _tabla(wsx, met.rename(columns={"index": "Muestra"}), 3, formatos={"tasa_malos": F_PCT, "gini": F_PCT,
                                                                              "ks": F_PCT, "auc": "0.000"})
    coef = pd.DataFrame({"Variable": ["Intercepto"] + [v.nombre for v in sc.variables], "Coeficiente": sc.beta,
                         "Error estándar": sc.ee, "Valor p": sc.pval})
    fila = _tabla(wsx, coef, fila + 2, formatos={"Coeficiente": "0.0000", "Error estándar": "0.0000", "Valor p": "0.0000"})
    wsx.cell(fila + 1, 1, f"Ajuste de calibración al intercepto: {sc.ajuste_calibracion:.4f}; PSI: {sc.metricas['psi']:.4f}").font = FUENTE
    fila = _tabla(wsx, sc.tabla_iv, fila + 3, formatos={"iv": "0.000"})
    fila = _tabla(wsx, sc.tarjeta, fila + 2, formatos={"participacion": F_PCT, "tasa_malos": F_PCT, "woe": "0.000"})
    fila = _tabla(wsx, sc.calificaciones.reset_index(), fila + 2,
                  formatos={"pd_modelo": "0.00%", "tasa_observada": "0.00%", "participacion": F_PCT, "pd_hasta": F_PCT})
    fila = _tabla(wsx, sc.cortes, fila + 2, formatos={"rechazo": F_PCT, "tasa_malos": F_PCT, "malos_evitados": F_PCT})
    _tabla(wsx, sc.psi_tabla, fila + 2, formatos={"desarrollo": F_PCT, "reciente": F_PCT, "contribucion": "0.0000"})

    # ---- fondeo con fórmulas
    wf = wb.create_sheet("Fondeo")
    _titulo(wf, "Base de préstamo y excedente financiero", "Fórmulas vivas: cambie los supuestos en la columna B.")
    filas = [
        ("Cartera bruta", f"=SUM({rng('C')})", F_MILES),
        ("Días máximos de mora elegibles", e["dias_elegible"], "0"),
        ("Cartera elegible", f"=SUMIF({rng('D')},\"<=\"&B5,{rng('C')})", F_MILES),
        ("Tasa de anticipo", e["anticipo"], F_PCT),
        ("Base de préstamo", "=B6*B7", F_MILES),
        ("Deuda con fondeadores", e["deuda"], F_MILES),
        ("Holgura", "=B8-B9", F_MILES),
        ("Cobertura de la deuda", "=B8/B9", F_DEC),
        ("", None, None),
        ("Tasa promedio de la cartera al día", float(foto["tasa_promedio"]), F_PCT),
        ("Saldo al día (mora <= 30)", f"=SUMIF({rng('D')},\"<=30\",{rng('C')})", F_MILES),
        ("Rendimiento", "=B13*B14/B4", F_PCT),
        ("Comisiones recurrentes", e["comisiones"], F_PCT),
        ("Patrimonio", e["patrimonio"], F_MILES),
        ("Costo promedio del fondeo", e["costo_fondeo"], F_PCT),
        ("Costo del fondeo (% cartera)", "=B18*MAX(B4-B17,0)/B4", F_PCT),
        ("Gastos operativos", e["opex"], F_PCT),
        ("Pérdida esperada 12m", float(fon["perdida_esperada"]), F_PCT),
        ("Excedente financiero", "=B15+B16-B19-B20-B21", F_PCT),
        ("Multiplicador de PD de equilibrio", "=1+B22/B21", F_DEC),
    ]
    for k, (t, v, fmt) in enumerate(filas, start=4):
        wf.cell(k, 1, t).font = NEGRITA if t in ("Base de préstamo", "Cobertura de la deuda", "Excedente financiero") else FUENTE
        if v is not None:
            c = wf.cell(k, 2, v)
            c.number_format, c.font = fmt, FUENTE
    wf.column_dimensions["A"].width = 40
    wf.column_dimensions["B"].width = 20
    hist = fon["historia"].reset_index()
    _tabla(wf, hist, 4, col=4, formatos={"cartera": F_MILES, "elegible": F_MILES, "base_prestamo": F_MILES,
                                         "pct_elegible": F_PCT, "mes": "yyyy-mm"})

    # ---- estrés
    we = wb.create_sheet("Estres")
    _titulo(we, "Pruebas de estrés y estrés inverso")
    es = r["estres"]
    fila = _tabla(we, es["tabla"], 3, formatos={"ecl": F_MILES, "ecl_adicional": F_MILES, "pct_patrimonio": F_PCT,
                                                 "saldo_etapa2": F_PCT, "icv30": F_PCT, "excedente": F_PCT,
                                                 "cobertura_deuda": F_DEC})
    inv = pd.DataFrame([es["inverso"]]).T.reset_index().rename(columns={"index": "Umbral", 0: "Multiplicador / valor"})
    fila = _tabla(we, inv, fila + 2, formatos={"Multiplicador / valor": "0.00"})
    _tabla(we, es["curva"], fila + 2, formatos={"ecl": F_MILES, "ecl_adicional": F_MILES, "pct_patrimonio": F_PCT})

    # ---- calidad y parámetros
    wq = wb.create_sheet("Calidad_datos")
    _titulo(wq, "Bitácora de calidad de datos")
    _tabla(wq, r["hallazgos"].tabla(), 3, anchos={"detalle": 70, "accion": 45, "hoja": 18})
    wpar = wb.create_sheet("Parametros")
    _titulo(wpar, "Parámetros de la ejecución y trazabilidad")
    tr = r["traza"]
    par = pd.DataFrame(list(tr["parametros"].items()), columns=["Parámetro", "Valor"])
    fila = _tabla(wpar, par, 3, anchos={"Parámetro": 45, "Valor": 50})
    arch = pd.DataFrame(list(tr["archivos"].items()), columns=["Archivo", "SHA-256"])
    _tabla(wpar, arch, fila + 2, anchos={"Archivo": 45, "SHA-256": 70})
    try:
        wb.save(ruta)
    except PermissionError:
        from .lectura import ErrorEntrada
        raise ErrorEntrada(f"No se pudo guardar {ruta.name}: ciérrelo en Excel y repita.")
    return ruta
